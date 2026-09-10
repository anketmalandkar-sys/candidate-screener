"""Unit tests for resume text extraction.

These sit directly on `extract_text(filename, data)` — the seam the upload route
delegates to — rather than reaching it through HTTP, so the format and encoding
edge cases are cheap to enumerate.
"""

from __future__ import annotations

import io
import zipfile

import pytest

from app.utilities.extract import (
    MAX_TEXT_CHARS,
    MAX_UPLOAD_BYTES,
    ExtractionError,
    extract_text,
)

pytestmark = pytest.mark.no_db


# --------------------------------------------------------------------------
# Non-text content behind a .txt name
# --------------------------------------------------------------------------


def test_binary_content_named_txt_is_rejected():
    """latin-1 will decode any byte string; a resume it is not.

    NUL bytes and a wall of control characters are the signal that a file
    claiming to be .txt is really binary, and the recruiter needs to be told
    that rather than handed mojibake and a score of zero.
    """
    data = bytes(range(8)) * 400  # \x00..\x07 repeated: no real text has this
    with pytest.raises(ExtractionError) as exc:
        extract_text("resume.txt", data)
    assert "text" in str(exc.value).lower()


def test_a_txt_with_a_few_control_characters_is_still_accepted():
    """The check is a ratio, not a tripwire — a stray form-feed is fine."""
    text = "Priya Nair\fSenior Engineer\nPython, Docker, Postgres\n"
    assert extract_text("resume.txt", text.encode()) == text


# --------------------------------------------------------------------------
# Encodings
# --------------------------------------------------------------------------


def test_utf16_txt_is_decoded():
    text = "Résumé — Priya Nair\nPython and Kubernetes\n"
    assert extract_text("resume.txt", text.encode("utf-16")) == text


def test_latin1_accented_txt_is_decoded():
    """utf-8 rejects a lone 0xE9; the latin-1 fallback must still read it.

    The odd length makes utf-16 fail cleanly first rather than decode to
    garbage, so this genuinely exercises the latin-1 branch.
    """
    original = "Jose - Go developer, cafe\xe9 x\n"
    assert extract_text("cv.txt", original.encode("latin-1")) == original


# --------------------------------------------------------------------------
# Extension handling
# --------------------------------------------------------------------------


def test_extension_match_is_case_insensitive():
    assert extract_text("RESUME.TXT", b"Python engineer") == "Python engineer"


def test_a_file_with_no_extension_is_rejected():
    with pytest.raises(ExtractionError) as exc:
        extract_text("resume", b"Python engineer")
    assert "Accepted formats" in str(exc.value)


def test_unknown_extension_is_rejected():
    with pytest.raises(ExtractionError) as exc:
        extract_text("resume.rtf", b"{\\rtf1 Python}")
    assert "Accepted formats" in str(exc.value)


# --------------------------------------------------------------------------
# Size and length bounds
# --------------------------------------------------------------------------


def test_empty_file_is_rejected():
    with pytest.raises(ExtractionError, match="empty"):
        extract_text("resume.txt", b"")


def test_a_file_at_the_size_limit_is_accepted():
    payload = b"Python " * (MAX_UPLOAD_BYTES // 7)
    assert len(payload) <= MAX_UPLOAD_BYTES
    assert extract_text("resume.txt", payload).startswith("Python")


def test_a_file_one_byte_over_the_limit_is_rejected():
    with pytest.raises(ExtractionError, match="larger than"):
        extract_text("resume.txt", b"a" * (MAX_UPLOAD_BYTES + 1))


def test_extracted_text_is_truncated_to_the_char_cap():
    payload = ("Python Docker Kubernetes " * 20_000).encode()
    result = extract_text("resume.txt", payload)
    assert len(result) == MAX_TEXT_CHARS


def test_whitespace_only_text_is_rejected_as_unreadable():
    with pytest.raises(ExtractionError) as exc:
        extract_text("resume.txt", b"   \n\t  \n   ")
    assert "couldn't read any text" in str(exc.value)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------


def _text_pdf(body: str) -> bytes:
    """A minimal single-page PDF carrying `body` as a real text object."""
    esc = body.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = b"BT /F1 12 Tf 72 720 Td (" + esc.encode("latin-1") + b") Tj ET"
    objs.append(
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream"
    )
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, obj in enumerate(objs, start=1):
        offsets.append(len(out))
        out += str(i).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n0000000000 65535 f \n"
    for off in offsets:
        out += ("%010d 00000 n \n" % off).encode()
    out += (
        b"trailer\n<< /Size "
        + str(len(objs) + 1).encode()
        + b" /Root 1 0 R >>\nstartxref\n"
        + str(xref_pos).encode()
        + b"\n%%EOF"
    )
    return bytes(out)


def test_text_bearing_pdf_is_extracted():
    pdf = _text_pdf("Priya Nair Python Docker Postgres")
    assert "Python Docker Postgres" in extract_text("resume.pdf", pdf)


def test_pdf_name_over_non_pdf_bytes_is_rejected_before_parsing():
    with pytest.raises(ExtractionError, match="not a PDF"):
        extract_text("resume.pdf", b"just some text, no %PDF- header")


def test_corrupt_pdf_body_gives_a_could_not_be_read_message():
    with pytest.raises(ExtractionError, match="could not be read"):
        extract_text("resume.pdf", b"%PDF-1.4\nthis is not a real pdf structure")


def test_encrypted_pdf_is_reported_as_password_protected():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("hunter2")
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(ExtractionError, match="password-protected"):
        extract_text("resume.pdf", buffer.getvalue())


def test_image_only_pdf_is_reported_as_scanned():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(ExtractionError, match="scanned"):
        extract_text("scan.pdf", buffer.getvalue())


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------


def _docx(paragraphs: list[str] = (), table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for line in paragraphs:
        document.add_paragraph(line)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_paragraphs_are_extracted():
    data = _docx(["Priya Nair", "Python, Docker and Postgres for six years"])
    text = extract_text("resume.docx", data)
    assert "Priya Nair" in text
    assert "Postgres" in text


def test_docx_table_cells_are_extracted():
    """Plenty of resumes are laid out entirely in a table; those cells must be
    read, or the candidate silently scores zero."""
    data = _docx(
        paragraphs=["Priya Nair"],
        table=[["Skill", "Years"], ["Python", "6"], ["Kubernetes", "3"]],
    )
    text = extract_text("resume.docx", data)
    assert "Python" in text
    assert "Kubernetes" in text


def test_docx_name_over_a_plain_zip_is_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "not a word document")

    with pytest.raises(ExtractionError, match="could not be read"):
        extract_text("resume.docx", buffer.getvalue())


def test_legacy_doc_named_docx_is_rejected_with_guidance():
    ole2_magic = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64
    with pytest.raises(ExtractionError, match="Older .doc files"):
        extract_text("resume.docx", ole2_magic)
