"""Resume text extraction from uploaded files.

Only the extracted TEXT is ever persisted — never the uploaded binary. That is
a deliberate scope cut: keeping the original file would pull in blob storage, a
retention policy, and PII-at-rest questions the brief does not ask about, and
the text is the only thing the scorer reads anyway.
"""

from __future__ import annotations

import io

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # 5 MB
MAX_TEXT_CHARS = 100_000

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

# Leading bytes we require for each format. The extension is a claim made by
# the client; the magic bytes are evidence. We check both.
PDF_MAGIC = b"%PDF-"
ZIP_MAGIC = b"PK\x03\x04"  # .docx is a zip container


class ExtractionError(Exception):
    """Raised with a message intended to be shown directly to the user."""


def _extension(filename: str) -> str:
    _, _, ext = filename.rpartition(".")
    return f".{ext.lower()}" if ext else ""


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError(
                "That PDF is password-protected. Remove the password, or paste "
                "the text instead."
            )
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, OSError) as exc:
        raise ExtractionError(f"That PDF could not be read ({exc}).") from exc


def _extract_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    except (ValueError, KeyError, OSError) as exc:
        raise ExtractionError(f"That Word file could not be read ({exc}).") from exc

    parts = [para.text for para in document.paragraphs]
    # Plenty of resumes lay everything out in a table; skipping tables would
    # silently produce an empty resume and a score of zero.
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


# Whitespace that is legitimately common in a plain-text resume. Anything else
# below U+0020, plus DEL, is a control character.
_TEXT_WHITESPACE = "\t\n\r\f\v"


def _looks_like_text(value: str) -> bool:
    """A cheap "is this a document, or a binary file behind a .txt name?" check.

    latin-1 decodes every possible byte string, so without this the encoding
    fallback below never actually fails — a JPEG renamed to .txt would sail
    through as mojibake and score zero with no explanation.
    """
    if "\x00" in value:
        return False
    if not value:
        return True
    control = sum(
        1
        for ch in value
        if (ord(ch) < 32 or ord(ch) == 127) and ch not in _TEXT_WHITESPACE
    )
    return control / len(value) <= 0.30


def _extract_txt(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        if not _looks_like_text(text):
            raise ExtractionError(
                "That file does not look like a text document — it may be a "
                "binary file with a .txt name. Paste the text instead."
            )
        return text
    raise ExtractionError("That text file is not in a recognised encoding.")


def extract_text(filename: str, data: bytes) -> str:
    """Validate an upload and return its plain text.

    Checks run cheapest-first: size, then claimed extension, then actual magic
    bytes, and only then the parse.
    """
    if not data:
        raise ExtractionError("That file is empty.")

    if len(data) > MAX_UPLOAD_BYTES:
        limit_mb = MAX_UPLOAD_BYTES // (1024 * 1024)
        raise ExtractionError(f"That file is larger than the {limit_mb} MB limit.")

    ext = _extension(filename or "")
    if ext not in ALLOWED_EXTENSIONS:
        allowed = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise ExtractionError(f"Unsupported file type. Accepted formats: {allowed}.")

    if ext == ".pdf":
        if not data.startswith(PDF_MAGIC):
            raise ExtractionError("That file is named .pdf but is not a PDF.")
        text = _extract_pdf(data)
    elif ext == ".docx":
        if not data.startswith(ZIP_MAGIC):
            raise ExtractionError(
                "That file is named .docx but is not a Word document. Older .doc "
                "files are not supported — save as .docx, or paste the text."
            )
        text = _extract_docx(data)
    else:
        text = _extract_txt(data)

    if not text or not text.strip():
        raise ExtractionError(
            "We couldn't read any text from that file — it may be a scanned "
            "image. Paste the text instead."
        )

    return text[:MAX_TEXT_CHARS]
